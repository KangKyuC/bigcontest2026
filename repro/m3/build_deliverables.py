# -*- coding: utf-8 -*-
"""Human-readable model card, evidence workbook, figures and portable parameters."""
from pathlib import Path
from datetime import datetime,timezone
import hashlib
import importlib.metadata
import json
import re
import zipfile
import joblib
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from m3lib import HERE,CFG,PRIMARY,dump,jsonable

R=HERE/'results';D=HERE/'data';M=HERE/'models';REPORT=HERE/'report';FIG=REPORT/'figures'
NAMES={'contact':'통화대상자','weekday_move':'평일 이동','holiday_move':'휴일 이동','flow_per_point':'관측점당 유동','flow_concentration':'유동 집중도','flow_change_log':'단기 유동 변화'}

def read(name):return pd.read_csv(R/name)

def markdown_table(df):
    cols=list(df.columns)
    def fmt(v):
        if pd.isna(v):return '—'
        if isinstance(v,(float,np.floating)):return f'{v:.3f}'
        return str(v).replace('|',' / ').replace('\n',' ')
    return '| '+' | '.join(map(str,cols))+' |\n|'+'|'.join(['---']*len(cols))+'|\n'+'\n'.join('| '+' | '.join(fmt(v) for v in row)+' |' for row in df.itertuples(index=False,name=None))

def figures(profile,summary,extra):
    FIG.mkdir(parents=True,exist_ok=True)
    from matplotlib.font_manager import fontManager
    available={f.name for f in fontManager.ttflist}
    chosen=next((f for f in ['Malgun Gothic','Noto Sans CJK KR','NanumGothic'] if f in available),'DejaVu Sans')
    plt.rcParams.update({'font.family':chosen,'axes.unicode_minus':False,'font.size':10,'axes.spines.top':False,'axes.spines.right':False})
    pdf=PdfPages(REPORT/'M3_핵심결과_도표.pdf')
    def finish(fig,name):
        fig.savefig(FIG/(name+'.png'),dpi=160,bbox_inches='tight');pdf.savefig(fig,bbox_inches='tight');plt.close(fig)
    cv=read('candidate_cv.csv')
    cv['label']=cv.candidate.replace({'shrink_1':'수축 Gaussian · 1분포','diag_1':'대각 Gaussian · 1분포','diag_2':'대각 GMM · 2유형','diag_3':'대각 GMM · 3유형','diag_4':'대각 GMM · 4유형','diag_5':'대각 GMM · 5유형'})
    fig,ax=plt.subplots(figsize=(10,5))
    y=np.arange(len(cv));ax.barh(y,cv['mean'],color=['#16817A' if s else '#9AAFC1' for s in cv.selected]);ax.errorbar(cv['mean'],y,xerr=cv.se,fmt='none',ecolor='#24384D',capsize=4)
    ax.set_yticks(y,cv.label);ax.set_xlabel('동 묶음 보류 로그밀도 평균 ± 1 SE (높을수록 보류자료 설명에 유리)')
    ax.set_title('여러 유형을 강제하지 않은 모델 비교');ax.text(0,-.18,'특성 벡터의 밀도 평가입니다. 고립 예측 정확도가 아닙니다.',transform=ax.transAxes,fontsize=9)
    fig.tight_layout();finish(fig,'01_model_comparison')
    cohorts=sorted(profile.cohort.unique(),key=lambda c:(c.split('_')[0],int(c.split('_')[1])))
    order=sorted(profile.dong.unique());fig,axs=plt.subplots(1,2,figsize=(15,9),sharey=True)
    for ax,field,title in zip(axs,['contact_seoul_peer_percentile','weekday_move_seoul_peer_percentile'],['연락: 같은 성·연령의 서울 다른 동 대비','평일 이동: 같은 성·연령의 서울 다른 동 대비']):
        a=profile.pivot(index='dong',columns='cohort',values=field).reindex(index=order,columns=cohorts)*100
        im=ax.imshow(a,aspect='auto',vmin=0,vmax=100,cmap='YlGnBu')
        ax.set_xticks(np.arange(len(cohorts)),[c.replace('_',' ').replace('60','60+') for c in cohorts],rotation=45,ha='right')
        ax.set_yticks(np.arange(len(order)),order);ax.set_title(title)
    fig.suptitle('강남 22동 × 성별2 × 성인연령5: 모든 220개 집단 프로필',fontsize=15)
    fig.subplots_adjust(bottom=.13,top=.91,wspace=.12,right=.86)
    cax=fig.add_axes([.89,.2,.018,.62]);fig.colorbar(im,cax=cax,label='비교위치 백분율 (위험도 아님)')
    finish(fig,'02_all_220_profiles')
    fig,ax=plt.subplots(figsize=(10,7))
    sc=ax.scatter(profile.contact_seoul_peer_percentile*100,profile.weekday_move_seoul_peer_percentile*100,
                  c=np.log1p(profile.flow_per_point),cmap='viridis',s=32,alpha=.8)
    for dong,sex,age in [('개포2동','여성',20),('논현1동','여성',30),('대치1동','여성',20)]:
        p=profile[(profile.dong==dong)&(profile.sex==sex)&(profile.age_floor==age)].iloc[0]
        ax.annotate(f'{dong} {age}대 {sex}',(p.contact_seoul_peer_percentile*100,p.weekday_move_seoul_peer_percentile*100),xytext=(5,9),textcoords='offset points',fontsize=9)
    for v in [25,75]:ax.axvline(v,color='#adb9c3',ls=':',lw=.7);ax.axhline(v,color='#adb9c3',ls=':',lw=.7)
    ax.set(xlim=(-3,103),ylim=(-3,103),xlabel='연락 지표의 서울 동년배 비교위치 (%)',ylabel='평일 이동 지표의 서울 동년배 비교위치 (%)',title='연락과 외출은 서로 다른 특성: 점 하나는 집계 집단')
    fig.colorbar(sc,ax=ax,label='log(1 + 공통 관측점당 유동)');fig.tight_layout();finish(fig,'03_contact_outings_place')
    metrics=read('contact_prediction_metrics.csv');fig,axs=plt.subplots(1,2,figsize=(12,5),sharey=True)
    models=['demographics','outings','outings_plus_SK'];labels=['성·연령','+ 외출','+ 외출 + SK']
    for ax,scheme,title in zip(axs,['dong','spatial'],['주 평가: 한 동씩 보류','보조 평가: 가까운 동 5블록 보류']):
        z=metrics[metrics.scheme.eq(scheme)].set_index('model').loc[models]
        bars=ax.bar(labels,z.mae,color=['#a5b6c8','#487eaa','#16817A'])
        ax.bar_label(bars,fmt='%.3f',padding=4);ax.set_title(title);ax.set_ylim(0,1.1);ax.set_ylabel('연락 지표 예측 MAE (작을수록 좋음)')
    fig.suptitle('SK 추가 기여는 검증 방식에 따라 달라짐',fontsize=15);fig.tight_layout();finish(fig,'04_SK_contribution')
    e=read('exploratory_gmm_bootstrap.csv');k=read('kmeans_bootstrap_stability.csv');fig,ax=plt.subplots(figsize=(9,5))
    ax.boxplot([e.oob_ari.dropna(),k.oob_ari.dropna()],tick_labels=['탐색 GMM 3유형','비교 K-means 2분할'],showfliers=False)
    ax.axhline(.75,color='#16817A',ls='--',lw=1,label='사전에 정한 중앙값 참고선 0.75')
    ax.set(ylabel='표집되지 않은 동의 배정 ARI',ylim=(-.1,1.05),title='동 재표집 200회: 안정적인 분할도 실제 주민 유형의 증명은 아님');ax.legend(loc='lower right')
    fig.tight_layout();finish(fig,'05_type_stability')
    pdf.close()

def evidence_cases(profile):
    specs=[('개포2동','여성',20,'연락 낮음·외출 높음'),('논현1동','여성',30,'연락 높음·외출 낮음'),('대치1동','여성',20,'세 행동 비교위치 모두 하위25%')]
    rows=[]
    for dong,sex,age,note in specs:
        r=profile[(profile.dong==dong)&(profile.sex==sex)&(profile.age_floor==age)].iloc[0]
        rows.append(dict(집단=f'{dong} {age}대 {sex}',연락지수=r.contact,연락_비교위치_pct=r.contact_seoul_peer_percentile*100,
                         평일이동_비교위치_pct=r.weekday_move_seoul_peer_percentile*100,휴일이동_비교위치_pct=r.holiday_move_seoul_peer_percentile*100,
                         관측점당유동=r.flow_per_point,해석=note))
    out=pd.DataFrame(rows);out.to_csv(R/'illustrative_cases.csv',index=False,encoding='utf-8-sig');return out

def card_context():
    # Optional descriptive card EDA, kept outside the six-feature model.
    path=R/'activity_monthly.csv'
    cols=['region','sex','age_floor','month','transactions','narrow_transactions','sports_transactions','gym_transactions','film_transactions']
    x=pd.read_csv(path) if path.is_file() else pd.DataFrame(columns=cols)
    if len(x):
        x=x[(x.region=='강남구')&(x.age_floor>=20)].copy()
    x['unit']='merchant district x sex x age x month; not residential dong'
    x.to_csv(R/'card_context_separate.csv',index=False,encoding='utf-8-sig')
    return dict(path='results/activity_monthly.csv',rows=len(x),role='optional separate context; never a model input')


def model_parameters(bundle):
    model=bundle['model']
    state={'version':'M3-1.0','candidate':bundle['candidate'],'type_adopted':bundle['adoption']['type_adopted'],'preprocessing':bundle['transformer'].state(),
           'features':PRIMARY,'scope':bundle['scope'],'training_input_sha256':bundle['input_sha256'],'interpretation':'distribution of six aggregate characteristics; not isolation probability'}
    if hasattr(model,'cov_'):state['density_model']=dict(mean=model.mean_,covariance=model.cov_,shrinkage=model.cov_model_.shrinkage_)
    else:state['density_model']=dict(weights=model.weights_,means=model.means_,covariances=model.covariances_,covariance_type=model.covariance_type)
    dump(state,M/'m3_parameters.json')

def workbook(profile,summary,extra):
    book=Workbook();book.remove(book.active)
    intro=pd.DataFrame([
        {'항목':'최종 모델','내용':'연속 프로필 + 수축 공분산 단일 Gaussian. 주민 위험유형은 채택하지 않음.'},
        {'항목':'관측 단위','내용':'강남22동×성별2×성인연령5=220집계셀. 2025.07~09. 사람220명이 아님.'},
        {'항목':'핵심 결과','내용':'GMM3유형 OOB ARI 중앙값0.390. K-means2분할은0.829이나 자연적 유형/고립의 검증은 아님.'},
        {'항목':'SK 추가 기여','내용':'동 보류 MAE는0.747→0.772. 5공간블록은0.835→0.732. 분할에 민감한 조건부 기여.'},
        {'항목':'확률과 결측','내용':'단일분포 소속확률은 의미가 없어 비움. 비어 있는 유형은 모든 사람이 같은 위험이라는 뜻이 아님.'},
        {'항목':'정책 모듈','내용':'지원 필요·정확한 나이·소득·참여 의사는 별도 입력. 집계 특성으로 자격을 판정하지 않음.'}])
    compact=profile[['dong','sex','age_floor','contact','contact_seoul_peer_percentile','weekday_move_seoul_peer_percentile','holiday_move_seoul_peer_percentile','flow_per_point','flow_concentration','flow_change_log','max_move_unestimated_share','quality_warning','type_adopted']].copy()
    for col in ['contact_seoul_peer_percentile','weekday_move_seoul_peer_percentile','holiday_move_seoul_peer_percentile','max_move_unestimated_share']:compact[col]*=100
    compact.columns=['행정동','성별','연령하한_60은60세이상','통화대상자_지수','연락_비교위치_pct','평일이동_비교위치_pct','휴일이동_비교위치_pct','공통관측점당유동','유동집중도_NxHHI','단기유동변화_log차이','원5세층최대미추정_pct','자료품질주의','위험유형채택여부']
    items={'읽는법':intro,'전체220프로필':compact,'사례':read('illustrative_cases.csv'),'모형선택':read('candidate_cv.csv'),
        '모형진단':read('candidate_fit_diagnostics.csv'),'동별보류선택':read('nested_dong_selection.csv'),'공간보류선택':read('nested_spatial_selection.csv'),
        '예측평가':read('contact_prediction_metrics.csv'),'동별예측오차':read('contact_prediction_fold_errors.csv'),'성연령별오차':read('contact_prediction_cohort_errors.csv'),
        '민감도':read('sensitivity.csv'),'유형안정성':pd.DataFrame([summary['exploratory_gmm_stability'],summary['kmeans_baseline_stability']]),
        '관측창비교':read('overlapping_window_comparison.csv'),'공간블록':read('spatial_blocks.csv'),
        '카드별도맥락':read('card_context_separate.csv'),'원본검사':pd.DataFrame(json.loads((D/'input_checks.json').read_text(encoding='utf-8'))),
        '검증체크':pd.DataFrame(json.loads((R/'validation.json').read_text(encoding='utf-8'))['checks'])}
    for ix,(name,frame) in enumerate(items.items(),1):
        ws=book.create_sheet(name);ws.append([str(c) for c in frame.columns])
        for row in frame.itertuples(index=False,name=None):
            values=[]
            for v in row:
                if isinstance(v,(dict,list)):values.append(json.dumps(v,ensure_ascii=False))
                elif pd.isna(v):values.append(None)
                else:values.append(jsonable(v))
            ws.append(values)
        ws.freeze_panes='D2' if name=='전체220프로필' else 'A2';ws.sheet_view.showGridLines=False
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='17324D');c.font=Font(name='맑은 고딕',bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=42
        for j,c in enumerate(frame.columns,1):ws.column_dimensions[get_column_letter(j)].width=22 if len(str(c))<25 else 32
        if name=='읽는법':ws.column_dimensions['B'].width=100
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font=Font(name='맑은 고딕',size=10);c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,float):c.number_format='0.000'
        if name=='읽는법':
            for n in range(2,ws.max_row+1):ws.row_dimensions[n].height=45
        table=Table(displayName='ModelEvidence'+str(ix),ref=f'A1:{get_column_letter(len(frame.columns))}{ws.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    book.save(REPORT/'M3_모델결과와_검증표.xlsx')

def report_text(summary,extra,profile,cases):
    cv=read('candidate_cv.csv')[['candidate','mean','se','selected']].rename(columns={'candidate':'모형','mean':'보류 로그밀도','se':'폴드 SE','selected':'선택'})
    perf=read('contact_prediction_metrics.csv')[['scheme','model','mae','rmse']]
    counts=len(profile);checks=json.loads((R/'validation.json').read_text(encoding='utf-8'))
    corr=json.loads((D/'prior_window_correction.json').read_text(encoding='utf-8'))
    s=summary['exploratory_gmm_stability'];k=summary['kmeans_baseline_stability'];ci=summary['SK_contact_mae_improvement']
    spatial=[]
    for a in extra['additional_partitions']:
        x=a['group_mean_improvement'];spatial.append({'공간블록수':a['blocks'],'외출기준MAE':a['mae_outings'],'SK추가MAE':a['mae_plus_SK'],'블록차이95%구간':f"{x['lower']:.3f} ~ {x['upper']:.3f}"})
    text=f'''# M3 v1.0 — 제작 결과와 통계 검증

완료일 2026-10-22. 실제 원본에서 재집계하고 모델을 학습·저장했다. 이 문서는 계획서가 아닌 실행 결과다.

## 1. 결과와 채택 결정

**강남 220개 집계 집단의 연락·외출·장소 이용 특성을 출력하는 M3를 구현했다. 최종 모형은 수축 공분산을 사용하는 단일 Gaussian과 연속 프로필이다. 고립 위험 유형은 채택하지 않았다.**

이번에 검토한 대각 공분산 GMM이 단일 연속분포보다 보류 자료를 더 잘 설명하지 못했고, 탐색 GMM3유형의 동 재표집 안정성도 부족했다. K-means2분할은 비교적 안정적이지만 안정적인 절단만으로 자연적인 주민 유형이 존재하거나 고립 위험을 분류했다는 주장은 성립하지 않는다. 이는 검토한 입력·후보·규칙 안의 결론이며 모든 가능한 군집모형에 대한 불가능성의 증명은 아니다. 군집을 반드시 만들어야 한다는 목표로 결과를 바꾸지 않았다.

이 결과는 M3 구현 실패와 다르다. 모델은 220셀 모두의 원값·동년배 비교위치·장소 활동환경·자료 품질·검토 질문을 재현 가능하게 출력한다. 다만 원래 구상한 '학습된 고립 위험 페르소나'까지 검증했다는 성과는 확보하지 못했다.

## 2. 실제 입력과 보정

- 강남22동 × 남녀2 × 성인5연령(20/30/40/50/60+) = {counts}셀. 사람 {counts}명이 아니며 독립적인 공간 단위는22동이다.
- 서울2025-10 통신정보의 최근3개월과 제공SK2025-07~09를 맞췄다. 서울2025-12는SK09~11과 맞췄다. 두 창은 겹치므로 미래 예측 평가가 아니다.
- 5세층은 동일가중 비교 지수로 합쳤다. 이는 전체 개인의 정확한 평균이 아니다. 인구가중 근사도 민감도로 비교했다.
- 강남을 제외한 서울402동(4,020개 성인 셀)을 연락·외출의 동년배 비교기준으로 사용했다. SK의 성연령별 중심과 척도는 학습 동에서만 적합했다.
- 핵심6변수: 통화대상자, 평일 이동, 휴일 이동, 공통 관측점당 유동, 관측점 유동 집중도, 창 시작~끝의 짧은 유동 변화다. 문자·근무·출근·관심집단은 주 입력에 넣지 않았다.
- SK는 원본CSV6개에서 공통15,178지점을 재구성하고 동별합계1,320개 성인 월 셀을 만들었다. 2026-10-22 원 실행에서는 기존 감사 합계와도 대조했다. 과거 결합표의12월 창에는7~9월 유동이 붙어 있었다. 새 입력에서는12월 창의 {corr['changed_december_cells']}셀을9~11월로 고쳤고, 원래 파일은 변경하지 않았다.
- 주 입력의 품질주의(원5세층 이동 미추정20% 초과)는2셀, 후속 창은27셀이다. 50%초과 배제조건에 해당하는 셀은 없어 두 창 모두220셀을 확보했다. 미추정 비율은 위험도가 아니다.

블록 크기 보정은 K-means 거리에서의 기여를 조정한다. GMM은 공분산도 추정하므로 같은 스케일링이 곧 블록별 동일 우도 가중을 뜻하지는 않는다. 평일·휴일 이동의 높은 상관 때문에 대각 단일 Gaussian만 기준으로 쓰면 군집의 우수성을 과대평가할 수 있다. 수축 전체 공분산의 단일 Gaussian을 별도 기준으로 둔 이유다.

## 3. 모형 선택과 공간 보류

같은 입력 공간에서 동을 묶은5-fold 평가와1-SE 단순화 규칙을 사용했다. 최종 선택은 수축 단일Gaussian이다. BIC·silhouette는 보조 진단이며 고립 분류 성능이 아니다. 입력 차원이 다른 모델의 BIC나 원시 우도를 직접 비교하지 않았다.

{markdown_table(cv)}

22회 한동보류의 각 학습셋 내부에서 후보를 다시 선택했다. 21회 수축 단일Gaussian, 1회 대각 단일Gaussian을 선택했고 다성분GMM은 선택되지 않았다. 가까운 동5블록 보류에서도 다성분이 선택되지 않았다. 공간 블록은 관측점 중심좌표의 근접 그룹으로, 공식 생활권이나 대표표본 권역은 아니다.

5개 폴드의 SE와1-SE는 모델 선택을 위한 근사적 규칙이다. 같은 데이터의 훈련영역이 겹치므로 독립 임상시험의 신뢰구간으로 해석하지 않는다. 모델의 로그밀도는 특성 벡터를 설명하는 수치이며 개인 고립 확률이나 위험점수가 아니다.

## 4. 군집 안정성을 따로 평가한 결과

동을 단위로200회 재표집하고, 매번 전처리와 모형을 다시 적합했다. 주 지표는 표집되지 않은 동(OOB)에 대한 배정 ARI다. 이름 정렬에는 학습 동만 사용했다.

| 비교 | OOB ARI 중앙값 | 하위10% | 해석 |
|---|---|---|---|
| 탐색 GMM3유형 | {s['oob_ari_median']:.3f} | {s['oob_ari_p10']:.3f} | 유형이 재표집에 민감해 확정 페르소나로 쓰기 어려움 |
| 비교 K-means2분할 | {k['oob_ari_median']:.3f} | {k['oob_ari_p10']:.3f} | 비교적 재현되는 분할. 고립 또는 자연적 주민 유형의 검증은 아님 |
| 선택 단일Gaussian | 해당없음 | 해당없음 | K=1의 ARI=1은 자명해 성능으로 보고하지 않음 |

단일Gaussian의 소속확률100%도 공개 프로필에서는 비운다. '모두 동일한 위험군'이라는 뜻이 아니다. OOB 로그밀도의5/50/95백분위는 측정값의 불확실성이 아니라 모델 재적합에 따른 분포 설명의 변동이며, 데이터 제공기관의 추정오차까지 포괄하지 않는다.

## 5. SK가 실제로 추가 정보를 주었는가

공통 목표를 통화대상자 지수로 고정했다. 성·연령만 / 성·연령+외출 / 성·연령+외출+SK를 같은 보류 셀에서 비교했다. 문자·SNS·통화량처럼 목표를 우회하는 연락 변수는 입력하지 않았다. Ridge 규제계수도 학습 내부 동 분할로 선택했다.

{markdown_table(perf)}

주 평가인 한동보류에서는 MAE가0.747에서0.772로 증가했다. 동별 짝지은 오차개선(양수=SK가 유리)은 {ci['mean']:.3f}, 근사95%구간은 [{ci['lower']:.3f}, {ci['upper']:.3f}]이다. **주 평가에서 SK의 일관된 예측 개선은 확인하지 못했다.**

가까운 동5블록 평가에서는0.835에서0.732로 감소했다. 이 결과만 선택해 증분가치가 검증됐다고 선언하지 않는다. 상반된 결과를 본 뒤 공간4/6블록을 추가 점검했으며, 주 평가를 대체하지 않는 탐색적 보조 결과다.

{markdown_table(pd.DataFrame(spatial))}

보조 공간분할에서는 개선 방향이 유지됐지만4/6블록 근사구간은0을 포함했다. 외출 기준모형의 동 평균 잔차에는 공간 유사성이 있었고, SK 추가 후 줄었다(Moran I 약0.467→-0.133). 이는 SK가 지역 맥락을 일부 설명할 가능성을 보여주는 탐색 진단이다. 블록수가4~6개로 적고 공간분할 의존성도 있어 보편적 개선이나 고립 탐지 성능을 입증한 것은 아니다.

보고서의 R²는 고립을 설명하는 비율이 아니다. 성·연령만의 기준모형도 높은 R²를 보여 원래 연령 차이의 영향이 크다. 발표의 중심 성능은 비교 가능한 MAE와 그 차이·불확실성으로 두는 것이 적절하다.

## 6. 실제로 더 구체화된 인사이트

아래 세 사례는 관측값을 설명하기 위해 선택했다. 위험 순위나 수혜 우선순위가 아니다. 비교위치는 서울 다른402동의 같은 성·연령 셀 가운데 위치한 백분율이다.

{markdown_table(cases)}

개포2동20대여성은 전화번호 연락 지표가 낮은 쪽에 있지만 외출 지표는 높은 쪽에 있다. 논현1동30대여성은 반대 방향이다. 따라서 외출이 많다는 이유로 연락도 활발하다고, 외출이 적다는 이유로 연락이 단절됐다고 단정할 수 없다. 같은 성·연령이어도 지역에서 확인할 질문이 달라진다.

대치1동20대여성 셀은 세 행동 비교위치가 모두 하위25%다. 이것은 그 셀의 평균에 대한 서술이며 그 안의 같은 개인이 세 조건을 동시에 가지거나 실제 지원이 필요하다는 증거가 아니다. 지역별 순위에 따라 주민을 고립자로 표시하지 않는다.

주 분석의220셀 중 연락이하위25%이면서 평일·휴일 이동이상위25%인 조합은11셀이다. 이 수치는 통신 행동의 불일치를 보여주는 집계 셀 수이며 주민 수가 아니다. 서로 다른 특성을 구별한다는 점이 기존의 단순 소비 감소보다 구체적인 산출물이다.

## 7. 민감도와 남은 불확실성

- 5세층 인구가중 근사와 동일가중의 공통 특성 순위는 유사했다(최저 Spearman 약0.988). 개인 평균의 분모가 검증됐다는 뜻은 아니다.
- 연령 보정을 제거하면 특성 순서가 크게 달라진다. 이는 같은 성·연령 내 비교라는 연구질문이 결과를 실질적으로 바꾼다는 뜻이다.
- 후속 관측창에서 연락·이동·유동 수준과 단기 변화의 안정성을 따로 읽어야 한다. 특히 짧은 유동 변화는 기간이 바뀌면 순서가 달라진다. 중첩창을 독립 미래 검증으로 제시하지 않는다.
- 스케일 변경이나 다른 변수를 제거해도 유지되는 개별 축의 순위1.0은 단조변환에서 예상되는 결과다. 이를 독립적인 통계적 강건성의 증명으로 세지 않는다.
- 22동의 공간 의존성, 통신사 관측 편향, 제공기관 추정오차, 집계에 따른 개인 내부 분포 손실은 남는다. 공개자료만으로 이를 제거했다고 주장하지 않는다.
- 정책 수요·실제 이용·고립 상태의 독립 정답은 이 모델에 없다. 조사 질문을 넘기는 것과 정책 의사결정의 실제 개선은 별도 검증 과제다.

## 8. 상원의 정책 분석에 넘기는 형식

policy_handoff.json은220개 집단 각각의 관측 특성과 확인할 질문을 제공한다. 지원 필요, 정확한 나이, 소득, 가구형태, 이용 가능 시간, 참여 의사는 미확인으로 남긴다. 관측값을 바꿔도 공식 정책의 자격요건을 변경하는 구조가 아니다.

상원이 같은 기능의 사업·운영 회차·자격·시간·장소·채널·대안을 연결하면, '이런 필요와 이용조건이 주어졌을 때 어떤 경로가 가능한가'를 검토할 수 있다. 이번 모델이 실제 지원 공백의 존재나 정책 보완 효과를 입증한 것은 아니다. 카드의 구×성×연령×월 맥락60행은 별도 파일로 제공하며 동별 학습 입력으로 복제하지 않았다.

## 9. 재현과 검수

공개 재현판 입력 검사 {len(json.loads((D/'input_checks.json').read_text(encoding='utf-8')))}개와 구현·통계설계 보호 검사 {len(checks['checks'])}개를 통과했다. 관측창 불일치·결측·다른 도시·10대 입력·중복 셀을 거절하고, 같은 동이 평가와 학습에 섞이지 않는지 확인했다. 변환이 원본을 수정하지 않는지, 저장 모델을 다시 읽었을 때 예측이 같은지, 탐색 GMM의 변수별 설명 합이 로그 소속비와 일치하는지도 검사했다.

주 파일:

- models/m3_bundle.joblib: 선택 모델과 전처리·범위·채택 상태.
- models/m3_parameters.json: 변환·평균·공분산의 읽을 수 있는 파라미터.
- results/group_profiles.csv: 전체220셀 프로필.
- results/policy_handoff.json: 정책팀 전달용 관측/가정 구분.
- report/M3_모델결과와_검증표.xlsx: 모델·입력·오차·민감도·전달표.
- report/M3_핵심결과_도표.pdf: 비교·전체집단·증분기여·안정성 그림.
- config.json, 분석계약.md, 보조진단_추가계획.md: 주 규칙과 결과 확인 후 보조 진단의 구별.
- prepare_data.py, train_evaluate.py, additional_diagnostics.py, predict.py, validate_model.py: 실행 코드.

이 패키지는 통계적으로 검토 가능한 M3의 첫 완성본이다. 통계적 완전성이나 대회 요구 충족을 보증한다는 뜻은 아니다. 현재 근거에 맞는 채택은 연속 프로필이며, 위험 유형이나 SK의 일관된 성능 향상을 주장하는 것은 보류한다.

## 10. 방법·자료 근거

- 서울시 행정동29개 통신정보: https://data.seoul.go.kr/dataList/OA-22267/F/1/datasetView.do
- 서울시민생활 설명서: https://data.seoul.go.kr/dataVisual/seoul/seoulLiving.do
- scikit-learn grouped/nested CV: https://scikit-learn.org/1.6/modules/cross_validation.html
- scikit-learn clustering: https://scikit-learn.org/1.6/modules/clustering.html
- scikit-learn covariance estimation: https://scikit-learn.org/1.6/modules/covariance.html
- 원본별 경로·해시: data/input_manifest.json. 기존 자료 확인 기록: reports/bigcontest_model_options_20261021/source_registry.json.
'''
    return text

def docx_report(text):
    doc=Document();s=doc.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.6);s.left_margin=Cm(1.5);s.right_margin=Cm(1.5)
    for name in ['Normal','Title','Heading 1','List Bullet']:
        st=doc.styles[name];st.font.name='맑은 고딕';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.color.rgb=RGBColor.from_string('17324D')
    def inline(p,t):
        for token in re.split(r'(\*\*[^*]+\*\*)',t):
            r=p.add_run(token[2:-2] if token.startswith('**') else token);r.bold=token.startswith('**')
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
            rows=[[v.strip() for v in ln.strip('|').split('|')] for ln in block]
            rows=[r for r in rows if not all(set(v)<=set('-: ') for v in r)]
            t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid'
            for row_no,row in enumerate(rows):
                cells=t.add_row().cells
                for j,val in enumerate(row):
                    p=cells[j].paragraphs[0];inline(p,val)
                    for rr in p.runs:rr.font.size=Pt(8 if len(row)>5 else 9);rr.bold=row_no==0
                trp=t.rows[-1]._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
                if row_no==0:trp.append(OxmlElement('w:tblHeader'))
        elif line.startswith('# '):inline(doc.add_paragraph(style='Title'),line[2:])
        elif line.startswith('## '):inline(doc.add_paragraph(style='Heading 1'),line[3:])
        elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
        else:inline(doc.add_paragraph(),line)
    doc.add_page_break();doc.add_heading('핵심 도표',level=1)
    for name in ['01_model_comparison','03_contact_outings_place','04_SK_contribution','05_type_stability']:
        doc.add_picture(str(FIG/(name+'.png')),width=Cm(17.5))
    footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run('M3 v1.0 · 2026-10-22 · ')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    doc.save(REPORT/'M3_제작결과와_통계검증.docx')


def main():
    REPORT.mkdir(exist_ok=True)
    summary=json.loads((R/'summary.json').read_text(encoding='utf-8'));extra=json.loads((R/'additional_diagnostics.json').read_text(encoding='utf-8'))
    profile=read('group_profiles.csv');bundle=joblib.load(M/'m3_bundle.joblib')
    cases=evidence_cases(profile);card=card_context();dump(card,R/'card_context_provenance.json')
    model_parameters(bundle);figures(profile,summary,extra);workbook(profile,summary,extra)
    text=report_text(summary,extra,profile,cases);(REPORT/'M3_제작결과와_통계검증.md').write_text(text,encoding='utf-8');docx_report(text)
    packages=['numpy','pandas','scipy','scikit-learn','joblib','threadpoolctl','openpyxl','matplotlib','python-docx']
    (REPORT/'environment_used.txt').write_text('\n'.join(n+'=='+importlib.metadata.version(n) for n in packages)+'\n',encoding='utf-8')
    out=[]
    for path in [REPORT/'M3_제작결과와_통계검증.docx',REPORT/'M3_모델결과와_검증표.xlsx',REPORT/'M3_핵심결과_도표.pdf',M/'m3_bundle.joblib',M/'m3_parameters.json',R/'group_profiles.csv',R/'policy_handoff.json']:
        out.append(dict(path=str(path.relative_to(HERE)),bytes=path.stat().st_size,sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    wb=load_workbook(REPORT/'M3_모델결과와_검증표.xlsx',read_only=True);sheet_counts={s.title:s.max_row-1 for s in wb.worksheets};wb.close()
    with zipfile.ZipFile(REPORT/'M3_제작결과와_통계검증.docx') as z:assert z.testzip() is None
    assert sheet_counts['전체220프로필']==220
    dump(dict(created_at_utc=datetime.now(timezone.utc).isoformat(),artifacts=out,sheets=sheet_counts,
              code_hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in HERE.glob('*.py')},
              content_verified=True,docx_render_verified=False),REPORT/'deliverable_manifest.json')
    print(json.dumps(dict(artifacts=len(out),sheets=sheet_counts,report=str(REPORT)),ensure_ascii=False))

if __name__=='__main__':main()
